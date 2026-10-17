# -*- coding: utf-8 -*-
"""一次性文档工具（P0 批次）：把两份介绍 PPT 的版本/用例口径统一到 v1.12.18。

背景：
  两份 PPT（docs/帮扶管理信息系统介绍.pptx 与根目录同名文件）在 v1.12.17 发布时
  已被 `update_pptx_versions_20261016.py` 更新到「v1.12.17 / 2026-10-17 /
  后端 12716 / 前端 6363」；本轮发布 v1.12.18 后需再对齐一次（版本与用例数）。

用法：
  cd backend && .venv313/Scripts/python.exe ../scripts/update_pptx_versions_20261017.py \\
      --backend-cases 12790 --frontend-cases 6363

行为：
  1. 先备份到 .scratch/w16-r28-fix/ppt-backup-<日期>/（已存在则不覆盖）；
  2. 文本框与表格单元格里把 1.12.17 / 12716 / 6363 / 2026-10-17 换成新值；
  3. 只替换上述**目标字面量**，历史版本（v1.12.10 / v1.12.5 …）不动。

为什么不用「按日期新建脚本」以外的方案：这是与上批同构的一次性工具（同一份 PPT、
同一条替换链），保留同样的形状最省复核成本；正式的 PPT 生成仍属人工交付物。
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
TARGETS = [
    ROOT / "docs" / "帮扶管理信息系统介绍.pptx",
    ROOT / "帮扶管理信息系统介绍.pptx",
]
BACKUP_DIR = ROOT / ".scratch" / "w16-r28-fix" / f"ppt-backup-{date.today().isoformat()}"

OLD_VERSION = "1.12.17"
NEW_VERSION = "1.12.18"
# 注意：PPT 里的用例数是**不带千分位**的裸数字（v1.12.17 那次替换后落成的形态），
# 这里必须按屏幕上的实际字面量写，写 "12,716" 会一处也匹配不到。
OLD_BACKEND_CASES = "12716"
OLD_FRONTEND_CASES = "6363"
# 前端口径的**测试文件数**（tech-stack 页的"299 个测试文件 · 6363 个用例"）：
# 用例数同为 6363 但文件数从 299 增至 301（同日全量实测），必须单独替换。
OLD_FRONTEND_FILES = "299 个测试文件"
NEW_FRONTEND_FILES = "301 个测试文件"
OLD_SOURCE_DATE = "2026-10-17"


def _replace_in_tf(text_frame, pairs) -> int:
    n = 0
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            text = run.text
            new = text
            for old, repl in pairs:
                new = new.replace(old, repl)
            if new != text:
                run.text = new
                n += 1
    return n


def replace_in_slide(slide, pairs) -> int:
    n = 0
    for shape in slide.shapes:
        if shape.has_text_frame:
            n += _replace_in_tf(shape.text_frame, pairs)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    n += _replace_in_tf(cell.text_frame, pairs)
    return n


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source-date", default=date.today().isoformat(),
                    help="替换「数据来源：<日期>」中的日期")
    ap.add_argument("--backend-cases", default=None, help="后端用例数（PPT 里的裸数字，如 12790）")
    ap.add_argument("--frontend-cases", default=None, help="前端用例数（PPT 里的裸数字，如 6363）")
    ap.add_argument("--frontend-files", default=NEW_FRONTEND_FILES,
                    help="前端测试文件数字面量（如 '301 个测试文件'；置空字符串可跳过）")
    args = ap.parse_args()

    pairs = [(OLD_VERSION, NEW_VERSION), (OLD_SOURCE_DATE, args.source_date)]
    if args.backend_cases:
        pairs.append((OLD_BACKEND_CASES, args.backend_cases))
    if args.frontend_cases:
        pairs.append((OLD_FRONTEND_CASES, args.frontend_cases))
    if args.frontend_files:
        pairs.append((OLD_FRONTEND_FILES, args.frontend_files))

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    for pptx in TARGETS:
        if not pptx.exists():
            print(f"[skip] 不存在: {pptx}")
            continue
        backup = BACKUP_DIR / pptx.name
        if not backup.exists():
            shutil.copy2(pptx, backup)

        prs = Presentation(str(pptx))
        replaced = 0
        for slide in prs.slides:
            replaced += replace_in_slide(slide, pairs)
        prs.save(str(pptx))
        print(f"{pptx.name}: 替换 {replaced} 处 run, 幻灯片 {len(prs.slides)} 页, 备份 {backup}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
