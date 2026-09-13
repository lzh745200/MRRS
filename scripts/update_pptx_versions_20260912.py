# -*- coding: utf-8 -*-
"""一次性文档工具：将 deliverables/*.pptx 版本号更新至 v1.12.6 并追加发布说明页。

用法：PYTHONPATH=<pptxenv> python scripts/update_pptx_versions_20260912.py
（先备份原文件到 deliverables/archive/ppt-backup-<date>/，再原地更新。）
"""
from __future__ import annotations

import re
import shutil
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PPT_DIR = ROOT / "deliverables"
BACKUP_DIR = PPT_DIR / f"archive/ppt-backup-{date.today().isoformat()}"

OLD_VERSIONS = ("v1.4.2", "1.4.2", "v1.7.2", "1.7.2")
NEW_VERSION = "v1.12.6"

RELEASE_BULLETS = [
    "遗留风险治理第一批：调度器 Timer 生命周期统一（可停止、无泄漏、分片清理接入调度）",
    "启动期不再执行 VACUUM / 全库扫描，消除启动窗口 database is locked 抖动",
    "multipart 请求体分级预检（备份恢复 10GB，其余 512MB），防超大请求打爆内存",
    "备份语义 fail-loud：一致性快照失败即中止，杜绝「陈旧但合法」的备份包",
    "安全加固延续：目录穿越、存储型 XSS、Windows 文件占用、并发竞态（v1.12.5）",
    "质量基线：后端 11089+ 用例 / 前端 302 文件全绿，双向覆盖率 100%",
]


def replace_versions_in_runs(prs: Presentation) -> int:
    """替换正文 run 中的旧版本号（保留原格式），返回替换次数。"""
    count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text = run.text
                    for old in OLD_VERSIONS:
                        if old in text:
                            text = text.replace(old, NEW_VERSION)
                    if text != run.text:
                        run.text = text
                        count += 1
    return count


def append_release_slide(prs: Presentation) -> None:
    layout = None
    for candidate in prs.slide_layouts:
        if candidate.name in ("标题和内容", "Title and Content", "1_Title and Content"):
            layout = candidate
            break
    if layout is None:
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]

    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.text = f"版本更新说明 · {NEW_VERSION}（2026-09-12）"

    body = None
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 1:
            body = shape
            break
    if body is None:
        left = Inches(0.6)
        top = Inches(1.8)
        width = Inches(8.5)
        height = Inches(4.6)
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for line in RELEASE_BULLETS:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.text = f"• {line}"
            p.font.size = Pt(16)
            first = False
        return

    tf = body.text_frame
    tf.clear()
    first = True
    for line in RELEASE_BULLETS:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(16)
        first = False


def main() -> None:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    for pptx in sorted(PPT_DIR.glob("*.pptx")):
        backup = BACKUP_DIR / pptx.name
        if not backup.exists():
            shutil.copy2(pptx, backup)

        prs = Presentation(str(pptx))
        replaced = replace_versions_in_runs(prs)
        before = len(prs.slides)
        append_release_slide(prs)
        prs.save(str(pptx))
        print(f"{pptx.name}: 替换版本号 {replaced} 处, 幻灯片 {before} -> {len(prs.slides)}")


if __name__ == "__main__":
    main()
