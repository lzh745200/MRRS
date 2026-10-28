"""统一表格读取层（WPS 兼容改造 P1，2026-10-26）。

改造前的三个真实问题（有代码证据）：

1. 导入端点与两个 service 直连 `openpyxl.load_workbook`，而 `.xls` 是 OLE2
   容器，openpyxl 根本不支持 —— 端点文案却承诺"支持 .xlsx/.xls"，用户按
   WPS「97-2003 格式」导出后必然失败；
2. WPS 的 `.et` 有 OOXML 与 OLE2 两种容器，只看扩展名判不出来；
3. 每处各写一套 try/except，报错文案五花八门，用户看不懂该怎么做。

本模块按**容器魔数**选解析器，并统一异常与文案：
    zip  → openpyxl
    ole2 → xlrd（解析 .xls 与旧版 WPS .et）
    其它 → UnsupportedSpreadsheetFormat（可读中文提示）

调用方新增/改造只需两件事：
    `rows = read_rows(content, filename)`                       —— 只关心二维数据
    `wb = load_workbook_any(content, filename)` + `wb.active` —— 需要 openpyxl 风格 API
"""

from __future__ import annotations

import logging
from typing import Any, List, Optional

from app.core.office_formats import (
    CONTAINER_OLE2,
    CONTAINER_ZIP,
    file_ext,
    sniff_container,
    spreadsheet_upload_hint,
    wps_save_as_message,
)

logger = logging.getLogger(__name__)


class UnsupportedSpreadsheetFormat(ValueError):
    """无法解析的表格文件。`user_message` 可直接回给前端（HTTP 400）。"""

    def __init__(self, message: str, *, filename: Optional[str] = None) -> None:
        super().__init__(message)
        self.user_message = message
        self.filename = filename


def _read_with_openpyxl(content: bytes) -> List[List[Any]]:
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
    try:
        return [list(row) for row in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _xlrd_cell(book, cell) -> Any:
    """把 xlrd 单元格值归一成 openpyxl 口径（空→None、整数→int、日期→datetime）。"""
    import xlrd

    ctype = cell.ctype
    value = cell.value
    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if ctype == xlrd.XL_CELL_TEXT:
        return value if value != "" else None
    if ctype == xlrd.XL_CELL_NUMBER:
        try:
            return int(value) if float(value).is_integer() else value
        except (TypeError, ValueError):
            return value
    if ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(value, book.datemode)
        except Exception as exc:  # pragma: no cover - 极端损坏文件  # noqa: BLE001 - 兜底但必须留痕
            logger.warning('_xlrd_cell 兜底失败（已降级）: %s', exc, exc_info=True)
            return value
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(value)
    return value


def _read_with_xlrd(content: bytes) -> List[List[Any]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        sheet = book.sheet_by_index(0)
        rows: List[List[Any]] = []
        for r in range(sheet.nrows):
            rows.append([_xlrd_cell(book, sheet.cell(r, c)) for c in range(sheet.ncols)])
        return rows
    finally:
        book.release_resources()


def _normalize_rows(rows: List[List[Any]]) -> List[List[Any]]:
    """补齐成矩形（不同解析器对短行处理不一，调用方按列下标取值才安全）。"""
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return rows
    return [list(r) + [None] * (width - len(r)) for r in rows]


def _attempt_order(container: str):
    if container == CONTAINER_ZIP:
        return (("openpyxl", _read_with_openpyxl), ("xlrd", _read_with_xlrd))
    if container == CONTAINER_OLE2:
        return (("xlrd", _read_with_xlrd), ("openpyxl", _read_with_openpyxl))
    return (("openpyxl", _read_with_openpyxl), ("xlrd", _read_with_xlrd))


def read_rows(content: bytes, filename: Optional[str] = None) -> List[List[Any]]:
    """按容器魔数读第一个工作表的全部行（空单元格为 None）。

    Raises:
        UnsupportedSpreadsheetFormat: 所有解析器都失败（文案含 WPS 指引）。
    """
    if not content:
        raise UnsupportedSpreadsheetFormat("文件为空，无法解析。", filename=filename)

    container = sniff_container(content)
    failures = []
    for engine, parser in _attempt_order(container):
        try:
            return _normalize_rows(parser(content))
        except ImportError as exc:  # 依赖缺失单独指出，避免误判为"文件坏了"
            failures.append(f"{engine}: 依赖缺失（{exc}）")
        except Exception as exc:  # noqa: BLE001 - 各解析器异常类型都不一致
            logger.warning('read_rows 兜底失败（已降级）: %s', exc, exc_info=True)
            failures.append(f"{engine}: {exc}")
        logger.debug("表格解析失败 engine=%s file=%s: %s", engine, filename, failures[-1])

    ext = file_ext(filename)
    if ext in {"wps", "wpt", "dps", "dpt"}:
        message = wps_save_as_message(filename)
    else:
        message = (
            f"{spreadsheet_upload_hint()}。"
            f"（解析失败原因：容器={container}；{'；'.join(failures[:2])}）"
        )
    raise UnsupportedSpreadsheetFormat(message, filename=filename)


# ── openpyxl 风格兼容层（供既有 `ws.iter_rows` 代码直接用）──


class _CellShim:
    __slots__ = ("value",)

    def __init__(self, value: Any) -> None:
        self.value = value


class _SheetShim:
    """最小工作表替身：只实现既有导入代码用到的 `iter_rows` / `max_row`。"""

    def __init__(self, rows: List[List[Any]]) -> None:
        self._rows = rows

    @property
    def max_row(self) -> int:
        return len(self._rows)

    @property
    def max_column(self) -> int:
        return max((len(r) for r in self._rows), default=0)

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: Optional[int] = None,
        values_only: bool = False,
        **_ignored: Any,
    ):
        start = max(min_row - 1, 0)
        end = len(self._rows) if max_row is None else min(max_row, len(self._rows))
        for row in self._rows[start:end]:
            if values_only:
                yield tuple(row)
            else:
                yield tuple(_CellShim(v) for v in row)


class _WorkbookShim:
    """最小工作簿替身：`active` / `sheetnames` / `close()` 与 openpyxl 同形。"""

    def __init__(self, rows: List[List[Any]], sheet_name: str = "Sheet1") -> None:
        self._sheet = _SheetShim(rows)
        self.sheetnames = [sheet_name]

    @property
    def active(self) -> _SheetShim:
        return self._sheet

    def __getitem__(self, key):  # 兼容 `wb[sheetname]` 取法（多工作表导入用）
        if key in self.sheetnames or key == 0:
            return self._sheet
        raise KeyError(key)

    def close(self) -> None:  # 与 openpyxl 接口对齐，无资源可释放
        return None


def load_workbook_any(content: bytes, filename: Optional[str] = None):
    """返回 openpyxl Workbook，或对 .xls/.et(OLE2) 返回等价的只读替身。

    调用方只需把 `openpyxl.load_workbook(BytesIO(x))` 换成
    `load_workbook_any(x, filename)`，其余 `wb.active.iter_rows(...)` 代码不动。
    """
    if sniff_container(content) == CONTAINER_ZIP:
        try:
            from io import BytesIO

            from openpyxl import load_workbook

            return load_workbook(filename=BytesIO(content), data_only=True)
        except Exception:  # noqa: BLE001 - 交由 read_rows 给出统一错误文案
            logger.debug("openpyxl 打开 zip 容器失败，回退通用读取：%s", filename)
    return _WorkbookShim(read_rows(content, filename))
