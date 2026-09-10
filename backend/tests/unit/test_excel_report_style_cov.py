"""覆盖率补充测试：`app/utils/excel_report_style.py`（他人新增模块，无专属测试）。

该模块由报表导出改造引入（统一军绿+金色 Excel 报表样式），被 `export_service` /
`report_service` / `supported_village_export_service` 复用，但此前没有专属测试，
导致全仓可覆盖集门禁（100%）差 13 行：

- `_coerce`：datetime 透传 / Decimal → float / Enum → value 三条分支；
- `make_subtitle`：带 unit 的分支；
- `apply_title_block`：带 system_name（多渲染一行系统名）的分支；
- `style_header_row`：未给 col_widths 时按表头长度自适应列宽的分支。

本文件只补测试，不改动被测源码。
"""

from datetime import date, datetime, time
from decimal import Decimal
from enum import Enum

import pytest
from openpyxl import Workbook

from app.utils import excel_report_style as style


class _Color(Enum):
    RED = "red"


class TestCoerce:
    @pytest.mark.parametrize("value", [None, True, 1, 1.5, "文本"])
    def test_passthrough_primitives(self, value):
        assert style._coerce(value) is value

    def test_datetime_like_passthrough(self):
        for value in (datetime(2026, 1, 1, 8, 30), date(2026, 1, 1), time(8, 30)):
            assert style._coerce(value) is value

    def test_decimal_to_float(self):
        assert style._coerce(Decimal("12.34")) == 12.34
        assert isinstance(style._coerce(Decimal("1")), float)

    def test_enum_to_value(self):
        assert style._coerce(_Color.RED) == "red"

    def test_other_objects_become_str(self):
        class _Thing:
            def __str__(self):
                return "thing"

        assert style._coerce(_Thing()) == "thing"
        assert style._coerce([1, 2]) == "[1, 2]"


class TestMakeSubtitle:
    def test_without_unit(self):
        text = style.make_subtitle(year=2026)
        assert "年度：2026" in text and "单位：" not in text and "生成时间：" in text

    def test_with_unit_and_extra(self):
        text = style.make_subtitle(unit="某部队", year=2026, extra="（内部）")
        assert "单位：某部队" in text
        assert "年度：2026" in text
        assert "（内部）" in text

    def test_empty(self):
        assert style.make_subtitle().startswith("生成时间：")


class TestApplyTitleBlock:
    def test_without_system_name(self):
        ws = Workbook().active
        start = style.apply_title_block(ws, "报表标题", subtitle="副标题", col_count=3)
        assert start == 5  # 主标题/副标题/装饰线/间隔 共 4 行
        assert ws.cell(row=1, column=1).value == "报表标题"

    def test_with_system_name_adds_row(self):
        ws = Workbook().active
        start = style.apply_title_block(
            ws, "报表标题", subtitle="副标题", col_count=3, system_name=style.SYSTEM_NAME
        )
        assert start == 6
        assert ws.cell(row=1, column=1).value == style.SYSTEM_NAME
        assert ws.cell(row=2, column=1).value == "报表标题"

    def test_col_count_zero_falls_back_to_one(self):
        ws = Workbook().active
        assert style.apply_title_block(ws, "T", col_count=0) == 5


class TestStyleHeaderRow:
    def test_adaptive_width_when_no_col_widths(self):
        ws = Workbook().active
        style.style_header_row(ws, 1, ["很长的表头名称", "短", "中等表头"])
        assert ws.cell(row=1, column=1).value == "很长的表头名称"
        assert ws.column_dimensions["A"].width >= 12
        assert ws.column_dimensions["B"].width == 12  # 短表头回落最小宽度

    def test_explicit_col_widths_are_capped(self):
        ws = Workbook().active
        style.style_header_row(ws, 2, ["a", "b"], col_widths=[30, 100])
        assert ws.column_dimensions["A"].width == 30
        assert ws.column_dimensions["B"].width == 45  # 上限 45

    def test_width_list_shorter_than_headers(self):
        ws = Workbook().active
        style.style_header_row(ws, 3, ["a", "b", "c"], col_widths=[20])
        assert ws.column_dimensions["A"].width == 20
        assert ws.column_dimensions["C"].width >= 12

    def test_decimal_and_enum_headers_are_coerced(self):
        ws = Workbook().active
        style.style_header_row(ws, 4, [Decimal("1.5"), _Color.RED])
        assert ws.cell(row=4, column=1).value == 1.5
        assert ws.cell(row=4, column=2).value == "red"


class TestStyleDataRegion:
    def test_empty_range_returns_early(self):
        ws = Workbook().active
        style.style_data_region(ws, 5, 4, 3)  # last_row < first_row

    def test_zebra_and_alignment(self):
        ws = Workbook().active
        for r in range(1, 4):
            for c in range(1, 3):
                ws.cell(row=r, column=c, value=f"{r}-{c}")
        style.style_data_region(ws, 1, 3, 2, align_map={1: "left", 2: "right"}, zebra=True)
        assert ws.cell(row=1, column=1).border.left.style == "thin"
        assert ws.cell(row=2, column=1).fill.start_color.rgb.endswith(style.MILITARY_ZEBRA)


class TestAutofitAndPrint:
    def test_autofit_uses_header_and_rows(self):
        ws = Workbook().active
        widths = style.autofit_widths(ws, ["名称", "说明"], [["甲", "很长的一段说明文字" * 3]])
        assert len(widths) == 2 and all(12 <= w <= 45 for w in widths)

    def test_setup_print_full(self):
        ws = Workbook().active
        style.setup_print(ws, col_count=4, last_row=20, header_row=5, footer_text="内部资料")
        assert ws.page_setup.paperSize == 9
        assert ws.print_title_rows == "$5:$5"  # openpyxl 规格化为绝对引用
        assert ws.freeze_panes == "A6"
        assert "内部资料" in ws.oddFooter.center.text

    def test_setup_print_minimal(self):
        ws = Workbook().active
        style.setup_print(ws, col_count=0, last_row=0, header_row=None, landscape=False, freeze=True)
        assert ws.page_setup.orientation == "portrait"
        assert ws.freeze_panes is None


class TestBuildReportSheet:
    def test_renders_new_sheet(self):
        wb = Workbook()
        ws, start, end = style.build_report_sheet(
            wb, "综合报表", ["名称", "金额"], [["甲", Decimal("1.50")], ["乙", None]],
            sheet_name="综合", subtitle="副标题", watermark="审计留痕",
        )
        assert ws.title == "综合"
        assert (start, end) == (6, 7)
        assert ws.cell(row=start, column=2).value == 1.5

    def test_empty_rows(self):
        wb = Workbook()
        ws, start, end = style.build_report_sheet(wb, "空报表", ["名称"])
        assert end == start - 1  # 无数据行时 data_end = header_row

    def test_reuses_existing_sheet_with_name(self):
        wb = Workbook()
        ws = wb.active
        out, _, _ = style.build_report_sheet(wb, "复用表", ["a"], ws=ws, sheet_name="改名")
        assert out is ws and ws.title == "改名"
