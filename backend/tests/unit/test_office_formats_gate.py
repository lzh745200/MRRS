"""办公文档格式门禁自身的用例（含负向：门禁必须能红）。

对应 `scripts/check_office_formats.py`：静态四条规则 + 运行时四类夹具自检。
负向用例用 tmp_path 造最小仓库结构，确保规则不是"恒绿"。
"""

import importlib.util
from io import BytesIO
from pathlib import Path

# backend/tests/unit/xxx.py → parents[3] = 仓库根
ROOT = Path(__file__).resolve().parents[3]
GATE_PATH = ROOT / "scripts" / "check_office_formats.py"


def _load_gate():
    spec = importlib.util.spec_from_file_location("check_office_formats_gate", GATE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


GATE = _load_gate()


def _write(tmp_path: Path, relative: str, content: str) -> Path:
    path = tmp_path / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path


def _rules(tmp_path: Path) -> set:
    return {item["rule"] for item in GATE.scan_static(tmp_path)}


def test_scan_static_flags_raw_openpyxl_import(tmp_path):
    _write(tmp_path, "backend/app/api/v1/foo.py", "from openpyxl import load_workbook\n\nwb = load_workbook(x)\n")

    assert "raw-openpyxl-import" in _rules(tmp_path)


def test_scan_static_flags_hardcoded_xls_whitelist(tmp_path):
    _write(
        tmp_path,
        "backend/app/api/v1/bar.py",
        'if not name.endswith((".xlsx", ".xls")):\n    raise ValueError\n',
    )

    assert "hardcoded-xls-whitelist" in _rules(tmp_path)


def test_scan_static_flags_non_wps_export_format(tmp_path):
    _write(
        tmp_path,
        "backend/app/api/v1/export.py",
        'headers = {"Content-Disposition": "attachment; filename=report.et"}\n',
    )

    assert "non-wps-export-format" in _rules(tmp_path)


def test_scan_static_accepts_standard_export_format(tmp_path):
    _write(
        tmp_path,
        "backend/app/api/v1/export.py",
        'headers = {"Content-Disposition": "attachment; filename=report.xlsx"}\n',
    )

    assert "non-wps-export-format" not in _rules(tmp_path)


def test_scan_static_flags_csv_without_bom(tmp_path):
    _write(
        tmp_path,
        "backend/app/api/v1/csv_export.py",
        "import csv\n\ndef build(rows):\n    buf = io.StringIO()\n    csv.writer(buf).writerows(rows)\n",
    )

    assert "csv-missing-bom" in _rules(tmp_path)


def test_scan_static_accepts_csv_with_bom(tmp_path):
    _write(
        tmp_path,
        "backend/app/api/v1/csv_export.py",
        "import csv\n\ndef build(rows):\n    buf = io.StringIO()\n    csv.writer(buf).writerows(rows)\n"
        "    return buf.getvalue().encode('utf-8-sig')\n",
    )

    assert "csv-missing-bom" not in _rules(tmp_path)


def test_scan_static_clean_repo_has_no_problems():
    assert GATE.scan_static(ROOT) == []


def test_runtime_fixtures_ok():
    """真依赖（openpyxl + xlrd）下，四类夹具（xlsx/xls/两种 .et 容器）都能读。"""
    assert GATE.check_runtime(ROOT) == []


def test_gate_main_passes_on_repo():
    assert GATE.main(["--root", str(ROOT)]) == 0


# ── 导出侧"WPS 能打开"的机器可证部分：xlsx 可回读、csv 带 BOM ──


def test_generated_xlsx_template_can_be_reopened():
    from openpyxl import load_workbook

    from app.services.excel_template_service import ExcelTemplateService

    content = ExcelTemplateService().generate_school_template()
    wb = load_workbook(BytesIO(content))
    try:
        assert wb.active is not None
        assert wb.active.max_row >= 1
    finally:
        wb.close()


def test_export_csv_has_utf8_bom():
    from app.api.v1.import_export.export import _to_csv

    content = _to_csv([{"名称": "帮扶村", "数量": 12}])

    assert content.startswith(b"\xef\xbb\xbf")
    assert content.decode("utf-8-sig").startswith("名称")
