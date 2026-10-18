"""w17-006 导出物「口径脚注」测试。

覆盖：
- export_service：口径脚注生成/格式化/时间换算/过滤条件整理/数据范围描述/
  导出人水印解析/Excel 表尾注入（含多列合并与单列不合并）
- 五类导出的实际接线（村庄/学校/项目/经费/综合报表）+ Word/PDF 文末脚注
- supported_village_export_service：xlsx 各 sheet 与 CSV 尾部脚注
"""

import datetime as dt
import zipfile
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.services.export_service import (
    ExcelExportService,
    _format_local_time,
    _format_scope_filters,
    apply_scope_footnote,
    build_scope_footnote,
    describe_data_scope,
    exporter_from_watermark,
    format_scope_footnote,
)


def _all_values(content: bytes) -> list:
    wb = load_workbook(BytesIO(content))
    values = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    values.append(str(cell))
    return values


def _contains(content: bytes, needle: str) -> bool:
    return any(needle in value for value in _all_values(content))


def _find_sheet_cell(content: bytes, needle: str):
    wb = load_workbook(BytesIO(content))
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.value is not None and needle in str(cell.value):
                    return wb[name], cell
    return None, None


# ══════════════════════════════════════════════════════════════════
# 基础函数
# ══════════════════════════════════════════════════════════════════


class TestFootnotePrimitives:
    def test_format_local_time_none(self):
        text = _format_local_time(None)
        assert len(text) == 16 and ":" in text

    def test_format_local_time_naive_treated_as_utc(self):
        naive = dt.datetime(2026, 1, 1, 12, 0)
        expected = naive.replace(tzinfo=dt.timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M")
        assert _format_local_time(naive) == expected

    def test_format_local_time_aware(self):
        aware = dt.datetime(2026, 1, 1, 12, 0, tzinfo=dt.timezone.utc)
        assert _format_local_time(aware) == aware.astimezone().strftime("%Y-%m-%d %H:%M")

    def test_format_scope_filters_none_and_empty_dict(self):
        assert _format_scope_filters(None) is None
        assert _format_scope_filters({}) is None

    def test_format_scope_filters_skips_empty_values(self):
        assert _format_scope_filters({"年度": None, "关键词": "", "区县": "   "}) is None

    def test_format_scope_filters_bool_and_values(self):
        text = _format_scope_filters({"年度": 2024, "是否振兴梯队": True, "含软删": False})
        assert text == "年度=2024；是否振兴梯队=是；含软删=否"

    def test_describe_data_scope_none(self):
        assert describe_data_scope(None) is None

    def test_describe_data_scope_all(self):
        user = SimpleNamespace(role="super_admin", is_superuser=True)
        assert describe_data_scope(user) == "全部数据（跨组织）"

    def test_describe_data_scope_own_dept(self):
        user = SimpleNamespace(role="admin", is_superuser=False, organization_id=5)
        assert describe_data_scope(user) == "本组织及下级组织"

    def test_describe_data_scope_own(self):
        user = SimpleNamespace(role="user", is_superuser=False, organization_id=5)
        assert describe_data_scope(user) == "仅本人创建的数据"

    def test_exporter_from_watermark_missing(self):
        assert exporter_from_watermark(None) is None
        assert exporter_from_watermark("") is None
        assert exporter_from_watermark("没有导出人标记") is None
        assert exporter_from_watermark("导出人：") is None

    def test_exporter_from_watermark_parsed(self):
        assert exporter_from_watermark("导出人：张三  2026-01-01 10:00") == "张三"

    def test_build_scope_footnote_full(self):
        aware = dt.datetime(2026, 1, 2, 3, 4, tzinfo=dt.timezone.utc)
        footnote = build_scope_footnote(
            data_cutoff=aware,
            filters={"年度": 2024, "空条件": None},
            include_deleted=True,
            scope="全部数据（跨组织）",
            generated_by="张三",
            generated_at=aware,
        )
        assert list(footnote) == ["数据截止", "过滤条件", "软删除", "数据范围", "生成人", "生成时间"]
        assert footnote["过滤条件"] == "年度=2024"
        assert footnote["软删除"] == "含软删除记录"
        assert footnote["数据范围"] == "全部数据（跨组织）"
        assert footnote["生成人"] == "张三"
        assert footnote["数据截止"] == footnote["生成时间"]

    def test_build_scope_footnote_defaults_omit_empty_fields(self):
        footnote = build_scope_footnote()
        assert list(footnote) == ["数据截止", "软删除", "生成时间"]
        assert footnote["软删除"] == "已排除软删除记录"

    def test_build_scope_footnote_generated_at_supplies_cutoff(self):
        aware = dt.datetime(2026, 5, 6, 7, 8, tzinfo=dt.timezone.utc)
        footnote = build_scope_footnote(generated_at=aware)
        assert footnote["数据截止"] == _format_local_time(aware)
        assert footnote["生成时间"] == _format_local_time(aware)

    def test_format_scope_footnote_empty(self):
        assert format_scope_footnote({}) == ""
        assert format_scope_footnote({"生成人": ""}) == ""

    def test_format_scope_footnote_text(self):
        text = format_scope_footnote({"数据截止": "2026-01-01 10:00", "软删除": "已排除软删除记录"})
        assert text == "口径说明：数据截止：2026-01-01 10:00；软删除：已排除软删除记录"


class TestApplyScopeFootnoteExcel:
    def test_single_column_no_merge(self):
        wb = Workbook()
        ws = wb.active
        row = apply_scope_footnote(ws, "口径说明：测试", start_row=3, col_count=1)
        assert row == 3
        assert ws.cell(row=3, column=1).value == "口径说明：测试"
        assert ws.row_dimensions[3].height == 18
        assert len(ws.merged_cells.ranges) == 0

    def test_multi_column_merged(self):
        wb = Workbook()
        ws = wb.active
        apply_scope_footnote(ws, "口径说明：测试", start_row=5, col_count=4)
        assert len(ws.merged_cells.ranges) == 1
        merged = list(ws.merged_cells.ranges)[0]
        assert (merged.min_col, merged.max_col) == (1, 4)
        assert ws.cell(row=5, column=1).font.color.rgb == "0064748B"


# ══════════════════════════════════════════════════════════════════
# ExcelExportService 五类导出接线
# ══════════════════════════════════════════════════════════════════


class TestExcelExportServiceFootnote:
    def test_five_exports_contain_footnote(self):
        svc = ExcelExportService()
        watermark = "导出人：张三  2026-01-01 10:00"
        exports = [
            svc.export_village_list([{"ID": 1, "名称": "甲村"}], watermark=watermark),
            svc.export_school_list([{"ID": 1, "名称": "甲校"}], watermark=watermark),
            svc.export_project_list([{"ID": 1, "名称": "甲项目"}], watermark=watermark),
            svc.export_fund_list([{"ID": 1, "名称": "甲经费"}], watermark=watermark),
            svc.export_comprehensive_report(
                {"总数": 1}, [{"ID": 1, "名称": "甲村"}], [{"ID": 1}], [{"ID": 1}], watermark=watermark
            ),
        ]
        for content in exports:
            assert _contains(content, "口径说明")
            assert _contains(content, "已排除软删除记录")
            assert _contains(content, "生成人：张三")

    def test_default_footnote_without_watermark_has_no_generator(self):
        svc = ExcelExportService()
        content = svc.export_village_list([{"ID": 1, "名称": "甲村"}])
        assert _contains(content, "口径说明")
        assert not _contains(content, "生成人：")

    def test_explicit_footnote_overrides_default(self):
        svc = ExcelExportService()
        content = svc.export_school_list(
            [{"ID": 1}], watermark="导出人：张三", footnote="口径说明：自定义脚注"
        )
        assert _contains(content, "自定义脚注")

    def test_comprehensive_each_sheet_gets_footnote(self):
        svc = ExcelExportService()
        content = svc.export_comprehensive_report(
            {"总数": 2},
            [{"ID": 1, "名称": "甲村"}],
            [{"ID": 1, "名称": "甲项目"}],
            [{"ID": 1, "名称": "甲经费"}],
            footnote="口径说明：统一脚注",
        )
        wb = load_workbook(BytesIO(content))
        for name in wb.sheetnames:
            values = [str(c) for row in wb[name].iter_rows(values_only=True) for c in row if c]
            assert any("统一脚注" in v for v in values), name

    def test_footnote_row_place_below_data(self):
        svc = ExcelExportService()
        content = svc.export_village_list(
            [{"ID": 1, "名称": "甲村"}, {"ID": 2, "名称": "乙村"}], footnote="口径说明：测试"
        )
        ws, cell = _find_sheet_cell(content, "口径说明：测试")
        # 数据 2 行（第 6-7 行）→ 空一行后脚注落在第 9 行
        assert cell.row == 9
        assert ws.max_row == 9


# ══════════════════════════════════════════════════════════════════
# report_export_service：Word / PDF 文末脚注
# ══════════════════════════════════════════════════════════════════


class TestReportFootnote:
    def test_report_footnote_uses_full_name_and_scope(self):
        from app.services.report_export_service import ReportExportService

        user = SimpleNamespace(role="admin", is_superuser=False, organization_id=5, full_name="李四")
        footnote = ReportExportService._report_footnote(user, 2024)
        assert footnote["过滤条件"] == "年度=2024"
        assert footnote["数据范围"] == "本组织及下级组织"
        assert footnote["生成人"] == "李四"

    def test_report_footnote_falls_back_to_username(self):
        from app.services.report_export_service import ReportExportService

        user = SimpleNamespace(role="user", is_superuser=False, username="operator")
        footnote = ReportExportService._report_footnote(user, 2024)
        assert footnote["生成人"] == "operator"
        assert footnote["数据范围"] == "仅本人创建的数据"

    def test_render_footnote_text_prefers_attached_dict(self):
        from app.services.report_export_service import ReportExportService

        attached = {"数据截止": "2026-01-01 10:00", "生成人": "王五"}
        text = ReportExportService._render_footnote_text({"footnote": attached})
        assert text == format_scope_footnote(attached)

    def test_render_footnote_text_falls_back_to_year(self):
        from app.services.report_export_service import ReportExportService

        text = ReportExportService._render_footnote_text({"year": 2024})
        assert "过滤条件：年度=2024" in text

    def test_render_footnote_text_without_year(self):
        from app.services.report_export_service import ReportExportService

        text = ReportExportService._render_footnote_text({})
        assert "过滤条件" not in text
        assert "口径说明" in text

    def test_word_contains_footnote(self):
        from app.services.report_export_service import ReportExportService

        data = {
            "year": 2024,
            "title": "测试报告",
            "sections": [{"title": "一、概况", "paragraphs": ["正文"], "table": None}],
            "footnote": {"数据截止": "2026-01-01 10:00", "生成人": "张三"},
        }
        content = ReportExportService().export_word("summary", data)
        with zipfile.ZipFile(BytesIO(content)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        assert "口径说明" in xml
        assert "生成人：张三" in xml

    def test_word_empty_footnote_dict_adds_nothing(self):
        from app.services.report_export_service import ReportExportService

        content = ReportExportService().export_word(
            "summary", {"year": 2024, "title": "t", "sections": [], "footnote": {}}
        )
        with zipfile.ZipFile(BytesIO(content)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        assert "口径说明" not in xml

    def test_pdf_contains_footnote_and_valid(self):
        from app.services.report_export_service import ReportExportService

        data = {
            "year": 2024,
            "title": "测试报告",
            "sections": [{"title": "一、概况", "paragraphs": ["正文"], "table": None}],
            "footnote": {"数据截止": "2026-01-01 10:00", "生成人": "张三"},
        }
        content = ReportExportService().export_pdf("summary", data)
        assert content[:5] == b"%PDF-"
        assert len(content) > 500

    def test_pdf_empty_footnote_dict_adds_nothing(self):
        from app.services.report_export_service import ReportExportService

        content = ReportExportService().export_pdf(
            "summary", {"year": 2024, "title": "t", "sections": [], "footnote": {}}
        )
        assert content[:5] == b"%PDF-"

    def test_fund_detail_exception_fallback_keeps_footnote(self):
        from app.services.report_export_service import ReportExportService

        db = MagicMock()
        db.query.side_effect = RuntimeError("db down")
        user = SimpleNamespace(role="super_admin", is_superuser=True, username="root")
        data = ReportExportService().generate_fund_detail_report_data(db, 2024, user)
        assert data["footnote"]["生成人"] == "root"
        assert any("查询失败" in p for p in data["sections"][0]["paragraphs"])

    def test_project_progress_exception_fallback_keeps_footnote(self):
        from app.services.report_export_service import ReportExportService

        db = MagicMock()
        db.query.side_effect = RuntimeError("db down")
        user = SimpleNamespace(role="super_admin", is_superuser=True, username="root")
        data = ReportExportService().generate_project_progress_report_data(db, 2024, user)
        assert data["footnote"]["过滤条件"] == "年度=2024"
        assert any("查询失败" in p for p in data["sections"][0]["paragraphs"])

    def test_generated_report_data_attaches_footnote(self, real_db_session):
        from app.services.report_export_service import ReportExportService

        user = SimpleNamespace(
            role="super_admin", is_superuser=True, organization_id=None, username="root"
        )
        data = ReportExportService().generate_village_summary_report_data(
            real_db_session, 2024, user
        )
        assert data["footnote"]["生成人"] == "root"
        assert data["footnote"]["过滤条件"] == "年度=2024"


# ══════════════════════════════════════════════════════════════════
# supported_village_export_service：xlsx / csv 脚注
# ══════════════════════════════════════════════════════════════════


class TestSupportedVillageExportFootnote:
    def _seed(self, db):
        from app.models.supported_village import SupportedVillage

        db.add(SupportedVillage(village_name="脚注村", county="测试县", department="某部"))
        db.commit()

    def test_generated_by_variants(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        assert SupportedVillageExportService(real_db_session)._generated_by() is None
        user = object()
        assert SupportedVillageExportService(real_db_session, current_user=user)._generated_by() is None
        only_name = SimpleNamespace(username="u1")
        assert (
            SupportedVillageExportService(real_db_session, current_user=only_name)._generated_by()
            == "u1"
        )
        both = SimpleNamespace(full_name="张三", username="u2")
        assert (
            SupportedVillageExportService(real_db_session, current_user=both)._generated_by()
            == "张三"
        )

    def test_build_footnote_filters(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        svc = SupportedVillageExportService(real_db_session)
        text = svc._build_footnote(
            year=2024,
            modules=["population", "income"],
            village_ids=[1, 2],
            department="某部",
            support_unit="某单位",
            tiered_level="示范级",
            keyword="甲",
            county="测试县",
            is_revitalization_tier=True,
        )
        assert "年度=2024" in text
        assert "帮扶村ID=1,2" in text
        assert "导出模块=人口数据、收入数据" in text
        assert "是否振兴梯队=是" in text
        assert "已排除软删除记录" in text

    def test_build_footnote_without_optional_filters(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        svc = SupportedVillageExportService(real_db_session)
        text = svc._build_footnote()
        assert "口径说明" in text
        assert "帮扶村ID" not in text
        assert "导出模块" not in text
        assert "生成人" not in text

    def test_export_xlsx_contains_footnote(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        self._seed(real_db_session)
        svc = SupportedVillageExportService(real_db_session)
        content, filename, stats = svc.export(
            year=2024, modules=["population"], keyword="脚注", county="测试县"
        )
        assert filename.endswith(".xlsx")
        assert _contains(content, "口径说明")
        assert _contains(content, "过滤条件：年度=2024")
        assert _contains(content, "关键词=脚注")

    def test_export_csv_contains_footnote(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        self._seed(real_db_session)
        svc = SupportedVillageExportService(real_db_session)
        content, filename, _stats = svc.export(format="csv", modules=["population"])
        assert filename.endswith(".csv")
        assert content.startswith(b"\xef\xbb\xbf")
        assert "# 口径说明" in content.decode("utf-8-sig")

    def test_export_csv_empty_data_keeps_only_footnote(self, real_db_session):
        from app.services.supported_village_export_service import SupportedVillageExportService

        svc = SupportedVillageExportService(real_db_session)
        content, _filename, _stats = svc.export(format="csv", modules=["population"], keyword="无")
        text = content.decode("utf-8-sig")
        assert "口径说明" in text
