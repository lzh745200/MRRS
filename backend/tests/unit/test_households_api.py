"""一户一档（W22 契约 §1/§2/§5）测试。

覆盖：
- 模型：三表字段、PII 透明加密（ORM 明文读写 / 落库密文）、软删语义；
- 服务：CRUD、筛选、数据域隔离、预警规则边界（0/90/91 天、大病/住房/突发困难）、
  风险记录登记/关闭、汇总回填 VillagePopulation（幂等 / 写回值 / 年度隔离）、
  导出数据组装与 xlsx/csv 字节；
- API：列表/详情/CRUD/成员/alerts/risk-records/rollup/export 全端点、404/409/400/422、
  软删与管理员 include_deleted、导出下载头与内容。

覆盖率目标：``app.models.household`` / ``app.schemas.household`` /
``app.services.household_service`` / ``app.api.v1.households`` 语句 + 分支 100%。
"""

from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient


def _user(**kw):
    base = dict(id=7, username="admin7", role="admin", organization_id=1, data_scope="all")
    base.update(kw)
    return SimpleNamespace(**base)


@pytest.fixture
def api_client():
    from fastapi import FastAPI
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from app.api.v1.households import router
    from app.core.database import get_db
    from app.core.exceptions import register_exception_handlers
    from app.core.security import get_current_user
    from app.models import Base
    from app.models.user import User  # noqa: F401  FK 目标表注册
    from app.models.organization import Organization  # noqa: F401
    from app.models.supported_village import (  # noqa: F401
        SupportedVillage,
        VillagePopulation,
    )
    from app.models.household import (  # noqa: F401
        Household,
        HouseholdMember,
        HouseholdRiskRecord,
    )
    from app.models.pairing import PairingRecord, VisitRecord  # noqa: F401  走访挂户表注册

    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(autocommit=False, autoflush=False, bind=engine)()
    state = {"user": _user()}

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[get_db] = lambda: session
    app.dependency_overrides[get_current_user] = lambda: state["user"]
    client = TestClient(app, raise_server_exceptions=False)
    try:
        yield client, session, state
    finally:
        app.dependency_overrides.clear()
        session.close()
        engine.dispose()


def _as(state, **kw):
    state["user"] = _user(**kw)


# ════════════════════════════════════════════════════════════════════
#  造数工具
# ════════════════════════════════════════════════════════════════════

def _mk_village(db, name="幸福村", org_id=1, active=True, **kw):
    from app.models.supported_village import SupportedVillage

    base = dict(
        village_name=name,
        organization_id=org_id,
        created_by=7,
        is_active=active,
        province="贵州省",
    )
    base.update(kw)
    row = SupportedVillage(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_household(db, village_id, **kw):
    from app.models.household import Household

    base = dict(
        supported_village_id=village_id,
        household_name="张三",
        household_type="general",
        organization_id=1,
        created_by=7,
    )
    base.update(kw)
    row = Household(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_member(db, household_id, **kw):
    from app.models.household import HouseholdMember

    base = dict(household_id=household_id, name="李四", organization_id=1, created_by=7)
    base.update(kw)
    row = HouseholdMember(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_risk_record(db, household_id, **kw):
    from app.models.household import HouseholdRiskRecord

    base = dict(
        household_id=household_id,
        risk_type="income_drop",
        risk_level="medium",
        status="open",
        organization_id=1,
        created_by=7,
    )
    base.update(kw)
    row = HouseholdRiskRecord(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_population(db, village_id, year, **kw):
    from app.models.supported_village import VillagePopulation

    row = VillagePopulation(supported_village_id=village_id, year=year, **kw)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_pairing(db, **kw):
    from app.models.pairing import PairingRecord

    base = dict(
        subject_type="unit",
        subject_name="某单位",
        partner_type="village",
        title="结对事项",
        status="active",
        organization_id=1,
        created_by=7,
        is_active=True,
    )
    base.update(kw)
    row = PairingRecord(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _mk_visit(db, pairing_id, **kw):
    from app.models.pairing import VisitRecord

    base = dict(
        pairing_id=pairing_id,
        visit_date=date(2026, 1, 1),
        visitors="甲",
        organization_id=1,
        created_by=7,
        is_active=True,
    )
    base.update(kw)
    row = VisitRecord(**base)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _post(client, payload, **kw):
    return client.post("/api/v1/households", json=payload, **kw)


def _base_payload(village_id, **kw):
    payload = {
        "supported_village_id": village_id,
        "household_name": "张三",
        "household_type": "general",
        "risk_level": "medium",
        "annual_income": 12000.5,
        "helper_name": "王干事",
        "helper_phone": "13800138000",
    }
    payload.update(kw)
    return payload


# ════════════════════════════════════════════════════════════════════
#  模型与 PII
# ════════════════════════════════════════════════════════════════════

class TestModelsAndPii:
    def test_pii_encrypted_at_rest_and_plaintext_via_orm(self, api_client):
        from sqlalchemy import text

        _client, db, _state = api_client
        village = _mk_village(db)
        h = _mk_household(db, village.id, helper_phone="13800138000")
        m = _mk_member(db, h.id, id_card="522122199001011234")

        # ORM 读取为明文（TypeDecorator 透明解密）
        db.expire_all()
        from app.models.household import Household, HouseholdMember

        assert db.query(Household).filter(Household.id == h.id).first().helper_phone == "13800138000"
        assert db.query(HouseholdMember).filter(HouseholdMember.id == m.id).first().id_card == (
            "522122199001011234"
        )
        # 落库为密文（enc.v1: 标记），证明 PII 走 EncryptedText 而非明文列
        raw_phone = db.execute(
            text("SELECT helper_phone FROM households WHERE id = :hid"), {"hid": h.id}
        ).scalar()
        raw_id_card = db.execute(
            text("SELECT id_card FROM household_members WHERE id = :mid"), {"mid": m.id}
        ).scalar()
        assert raw_phone.startswith("enc.v1:")
        assert raw_id_card.startswith("enc.v1:")

    def test_model_to_dict_and_is_deleted(self, api_client):
        _client, db, _state = api_client
        village = _mk_village(db)
        h = _mk_household(
            db,
            village.id,
            household_code="H-001",
            annual_income=5000,
            has_serious_illness=True,
            housing_safe=False,
            last_visit_date=date(2026, 1, 1),
        )
        data = h.to_dict()
        assert data["household_code"] == "H-001"
        assert data["annual_income"] == 5000.0
        assert data["has_serious_illness"] is True
        assert data["housing_safe"] is False
        assert data["last_visit_date"] == "2026-01-01"
        assert data["is_deleted"] is False
        assert data["created_at"]
        h.is_active = False
        assert h.is_deleted is True

        m = _mk_member(db, h.id, income=100.5)
        mdata = m.to_dict()
        assert mdata["income"] == 100.5
        assert mdata["is_deleted"] is False

        r = _mk_risk_record(db, h.id, reason="收入骤减")
        rdata = r.to_dict()
        assert rdata["reason"] == "收入骤减"
        assert rdata["detected_at"]
        assert rdata["status"] == "open"

    def test_model_repr(self, api_client):
        _client, db, _state = api_client
        village = _mk_village(db)
        h = _mk_household(db, village.id)
        assert "Household" in repr(h)
        m = _mk_member(db, h.id)
        assert "HouseholdMember" in repr(m)
        r = _mk_risk_record(db, h.id)
        assert "HouseholdRiskRecord" in repr(r)


# ════════════════════════════════════════════════════════════════════
#  Schema 校验（422）
# ════════════════════════════════════════════════════════════════════

class TestSchemaValidation:
    @pytest.mark.parametrize(
        "overrides",
        [
            {"household_name": ""},
            {"annual_income": -1},
            {"household_type": "unknown"},
            {"risk_type": "unknown"},
            {"risk_level": "fatal"},
            {"helper_phone": "12345"},
            {"migrant_workers": -2},
            {"supported_village_id": 0},
        ],
    )
    def test_create_validation_errors(self, api_client, overrides):
        client, db, _state = api_client
        village = _mk_village(db)
        resp = _post(client, _base_payload(village.id, **overrides))
        assert resp.status_code == 422

    def test_phone_empty_and_blank_normalized(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        resp = _post(client, _base_payload(village.id, helper_phone=""))
        assert resp.status_code == 200
        assert resp.json()["data"]["helper_phone"] is None
        resp = _post(client, _base_payload(village.id, helper_phone="   ", household_name="李四"))
        assert resp.status_code == 200
        assert resp.json()["data"]["helper_phone"] is None

    def test_update_phone_invalid_and_valid(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = _post(client, _base_payload(village.id)).json()["data"]["id"]
        bad = client.put(f"/api/v1/households/{hid}", json={"helper_phone": "abc"})
        assert bad.status_code == 422
        good = client.put(f"/api/v1/households/{hid}", json={"helper_phone": "13900139000"})
        assert good.status_code == 200
        assert good.json()["data"]["helper_phone"] == "13900139000"

    def test_member_validation_errors(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = _post(client, _base_payload(village.id)).json()["data"]["id"]
        assert client.post(f"/api/v1/households/{hid}/members", json={"name": ""}).status_code == 422
        assert (
            client.post(f"/api/v1/households/{hid}/members", json={"name": "甲", "income": -1}).status_code
            == 422
        )
        assert (
            client.post(
                f"/api/v1/households/{hid}/members", json={"name": "甲", "birth_year": 1700}
            ).status_code
            == 422
        )

    def test_risk_record_validation_errors(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = _post(client, _base_payload(village.id)).json()["data"]["id"]
        resp = client.post(
            f"/api/v1/households/{hid}/risk-records", json={"risk_type": "unknown"}
        )
        assert resp.status_code == 422


# ════════════════════════════════════════════════════════════════════
#  户 CRUD / 列表 / 详情
# ════════════════════════════════════════════════════════════════════

class TestHouseholdCrud:
    def test_create_and_detail(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db, name="幸福村")
        resp = _post(client, _base_payload(village.id, household_code="H-001"))
        assert resp.status_code == 200
        body = resp.json()
        assert body["code"] == 200
        data = body["data"]
        assert data["household_name"] == "张三"
        assert data["household_type"] == "general"
        assert data["helper_phone"] == "13800138000"
        assert data["organization_id"] == 1
        assert data["created_by"] == 7
        assert data["is_deleted"] is False

        detail = client.get(f"/api/v1/households/{data['id']}").json()["data"]
        assert detail["village_name"] == "幸福村"
        assert detail["member_count"] == 0
        assert detail["risk_alerts"] == []
        assert detail["members"] == []
        assert detail["risk_records"] == []
        assert detail["visits"] == []  # 无走访挂户记录

    def test_create_village_not_found_and_cross_scope(self, api_client):
        client, db, state = api_client
        _mk_village(db, name="外村", org_id=2)
        assert _post(client, _base_payload(9999)).status_code == 404
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert client.get("/api/v1/households").status_code == 200
        _as(state)
        village = _mk_village(db)
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert _post(client, _base_payload(village.id)).status_code == 404

    def test_create_duplicate_household_code_conflict(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        assert _post(client, _base_payload(village.id, household_code="H-001")).status_code == 200
        dup = _post(client, _base_payload(village.id, household_code="H-001", household_name="李四"))
        assert dup.status_code == 409
        assert "户编号" in dup.json()["message"]
        # 空编号不参与唯一性校验
        assert _post(client, _base_payload(village.id)).status_code == 200
        assert _post(client, _base_payload(village.id, household_name="王五")).status_code == 200

    def test_list_envelope_filters_and_enrichment(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db, name="幸福村")
        h1 = _mk_household(
            db, village.id, household_name="张三", household_type="poverty", risk_level="high"
        )
        h2 = _mk_household(
            db, village.id, household_name="李四", household_type="monitor", risk_level="low"
        )
        _mk_member(db, h1.id)
        _mk_member(db, h1.id)
        resp = client.get("/api/v1/households")
        assert resp.status_code == 200
        body = resp.json()
        assert body["code"] == 200
        data = body["data"]
        assert data["total"] == 2
        assert data["page"] == 1
        items = {item["household_name"]: item for item in data["items"]}
        assert items["张三"]["village_name"] == "幸福村"
        assert items["张三"]["member_count"] == 2
        assert items["李四"]["member_count"] == 0

        assert client.get("/api/v1/households", params={"village_id": village.id}).json()[
            "data"
        ]["total"] == 2
        assert client.get("/api/v1/households", params={"village_id": 9999}).json()["data"][
            "total"
        ] == 0
        assert client.get("/api/v1/households", params={"risk_level": "high"}).json()["data"][
            "total"
        ] == 1
        assert client.get("/api/v1/households", params={"household_type": "monitor"}).json()[
            "data"
        ]["total"] == 1
        assert client.get("/api/v1/households", params={"keyword": "张"}).json()["data"][
            "total"
        ] == 1
        h2.helper_name = "赵干事"
        h2.household_code = "H-009"
        db.commit()
        assert client.get("/api/v1/households", params={"keyword": "赵干事"}).json()["data"][
            "total"
        ] == 1
        assert client.get("/api/v1/households", params={"keyword": "H-009"}).json()["data"][
            "total"
        ] == 1
        assert client.get("/api/v1/households", params={"keyword": "不存在"}).json()["data"][
            "total"
        ] == 0

    def test_list_pagination(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        for index in range(3):
            _mk_household(db, village.id, household_name=f"户{index}")
        data = client.get("/api/v1/households", params={"page": 2, "page_size": 2}).json()["data"]
        assert data["total"] == 3
        assert data["page"] == 2
        assert len(data["items"]) == 1

    def test_list_empty(self, api_client):
        client, _db, _state = api_client
        data = client.get("/api/v1/households").json()["data"]
        assert data == {"items": [], "total": 0, "page": 1, "page_size": 20}

    def test_update_partial_and_village_change(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        village2 = _mk_village(db, name="光明村")
        hid = _post(client, _base_payload(village.id, household_code="H-001")).json()["data"]["id"]

        # 只改姓名：不触发村/编号校验
        resp = client.put(f"/api/v1/households/{hid}", json={"household_name": "李四"})
        assert resp.status_code == 200
        assert resp.json()["data"]["household_name"] == "李四"
        assert resp.json()["data"]["household_code"] == "H-001"

        # 同编号在自身更新时不算冲突
        resp = client.put(f"/api/v1/households/{hid}", json={"household_code": "H-001"})
        assert resp.status_code == 200

        # 换村并改编号：目标村可用
        resp = client.put(
            f"/api/v1/households/{hid}",
            json={"supported_village_id": village2.id, "household_code": "H-002"},
        )
        assert resp.status_code == 200
        assert resp.json()["data"]["supported_village_id"] == village2.id

    def test_update_village_not_found_and_duplicate_code(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        _mk_household(db, village.id, household_name="李四", household_code="H-002")
        hid = _post(client, _base_payload(village.id, household_code="H-001")).json()["data"]["id"]
        assert (
            client.put(
                f"/api/v1/households/{hid}", json={"supported_village_id": 9999}
            ).status_code
            == 404
        )
        dup = client.put(f"/api/v1/households/{hid}", json={"household_code": "H-002"})
        assert dup.status_code == 409

    def test_detail_not_found_and_cross_scope(self, api_client):
        client, db, state = api_client
        village = _mk_village(db, org_id=1)
        h = _mk_household(db, village.id)
        assert client.get("/api/v1/households/9999").status_code == 404
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert client.get(f"/api/v1/households/{h.id}").status_code == 404

    def test_detail_with_members_and_risk_records(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        h = _mk_household(db, village.id, household_type="monitor", last_visit_date=None)
        _mk_member(db, h.id, name="成员甲")
        _mk_risk_record(db, h.id, risk_type="income_drop", reason="收入骤减")
        data = client.get(f"/api/v1/households/{h.id}").json()["data"]
        assert [m["name"] for m in data["members"]] == ["成员甲"]
        assert len(data["risk_records"]) == 1
        assert data["risk_records"][0]["risk_type"] == "income_drop"
        # 监测户 + 从未走访 → 命中 medium 级连续未走访预警
        assert [alert["code"] for alert in data["risk_alerts"]] == ["no_visit"]
        assert data["visits"] == []

    def test_detail_includes_visits(self, api_client):
        """走访挂户（W22）：详情返回该户有效走访，按 visit_date 倒序且字段固定。"""
        client, db, _state = api_client
        village = _mk_village(db)
        household = _mk_household(db, village.id, household_name="张三")
        other_household = _mk_household(db, village.id, household_name="李四")
        pairing = _mk_pairing(db, village_id=village.id)
        earlier = _mk_visit(
            db,
            pairing.id,
            household_id=household.id,
            visit_date=date(2026, 1, 1),
            visitors="甲、乙",
            purpose="春节慰问",
            content="送去米面油",
            materials="米面油各一份",
        )
        later = _mk_visit(
            db,
            pairing.id,
            household_id=household.id,
            visit_date=date(2026, 3, 1),
            visitors="丙",
        )
        # 软删走访 / 其它户走访都不应返回
        _mk_visit(
            db,
            pairing.id,
            household_id=household.id,
            visit_date=date(2026, 2, 1),
            visitors="丁",
            is_active=False,
        )
        _mk_visit(
            db,
            pairing.id,
            household_id=other_household.id,
            visit_date=date(2026, 4, 1),
            visitors="戊",
        )

        data = client.get(f"/api/v1/households/{household.id}").json()["data"]
        visits = data["visits"]
        assert [visit["id"] for visit in visits] == [later.id, earlier.id]
        assert set(visits[0]) == {
            "id",
            "pairing_id",
            "visit_date",
            "visitors",
            "purpose",
            "content",
            "materials",
        }
        assert visits[0]["pairing_id"] == pairing.id
        assert visits[0]["visit_date"] == "2026-03-01"
        assert visits[1]["visit_date"] == "2026-01-01"
        assert visits[1]["visitors"] == "甲、乙"
        assert visits[1]["purpose"] == "春节慰问"
        assert visits[1]["content"] == "送去米面油"
        assert visits[1]["materials"] == "米面油各一份"

    def test_delete_soft_and_include_deleted(self, api_client):
        client, db, state = api_client
        village = _mk_village(db)
        hid = _post(client, _base_payload(village.id)).json()["data"]["id"]
        assert client.delete(f"/api/v1/households/{hid}").status_code == 200
        # 默认隐藏
        assert client.get("/api/v1/households").json()["data"]["total"] == 0
        assert client.get(f"/api/v1/households/{hid}").status_code == 404
        # 管理员 include_deleted 可见
        data = client.get(
            "/api/v1/households", params={"include_deleted": "true"}
        ).json()["data"]
        assert data["total"] == 1
        assert data["items"][0]["is_deleted"] is True
        detail = client.get(
            f"/api/v1/households/{hid}", params={"include_deleted": "true"}
        )
        assert detail.status_code == 200
        # 非管理员被静默降级
        _as(state, id=9, role="user", organization_id=1, data_scope="org")
        data = client.get(
            "/api/v1/households", params={"include_deleted": "true"}
        ).json()["data"]
        assert data["total"] == 0
        assert (
            client.get(
                f"/api/v1/households/{hid}", params={"include_deleted": "true"}
            ).status_code
            == 404
        )
        # 软删后不可更新/重复删除
        _as(state)
        assert client.put(f"/api/v1/households/{hid}", json={"household_name": "x"}).status_code == 404
        assert client.delete(f"/api/v1/households/{hid}").status_code == 404

    def test_list_data_scope_isolation(self, api_client):
        client, db, state = api_client
        village1 = _mk_village(db, name="本村", org_id=1)
        village2 = _mk_village(db, name="他村", org_id=2)
        _mk_household(db, village1.id, household_name="本户")
        other = _mk_household(db, village2.id, household_name="他户", organization_id=2)
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        data = client.get("/api/v1/households").json()["data"]
        assert data["total"] == 1
        assert data["items"][0]["household_name"] == "他户"
        assert client.get(f"/api/v1/households/{other.id}").status_code == 200


# ════════════════════════════════════════════════════════════════════
#  家庭成员
# ════════════════════════════════════════════════════════════════════

class TestMembers:
    def _create_household(self, client, village_id):
        return _post(client, _base_payload(village_id)).json()["data"]["id"]

    def test_member_crud(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = self._create_household(client, village.id)

        created = client.post(
            f"/api/v1/households/{hid}/members",
            json={
                "name": "成员甲",
                "relation": "配偶",
                "id_card": "522122199001011234",
                "gender": "女",
                "birth_year": 1990,
                "education": "初中",
                "health_status": "健康",
                "employment_status": "务工",
                "income": 8000,
                "is_special_support": True,
            },
        )
        assert created.status_code == 200
        member = created.json()["data"]
        assert member["name"] == "成员甲"
        assert member["organization_id"] == 1
        assert member["created_by"] == 7
        assert member["id_card"] == "522122199001011234"

        listed = client.get(f"/api/v1/households/{hid}/members").json()["data"]
        assert listed["total"] == 1
        assert listed["items"][0]["name"] == "成员甲"

        updated = client.put(
            f"/api/v1/households/{hid}/members/{member['id']}", json={"income": 9000.5}
        )
        assert updated.status_code == 200
        assert updated.json()["data"]["income"] == 9000.5

        deleted = client.delete(f"/api/v1/households/{hid}/members/{member['id']}")
        assert deleted.status_code == 200
        listed = client.get(f"/api/v1/households/{hid}/members").json()["data"]
        assert listed["total"] == 0
        # 软删成员不再计入列表富化的成员数
        item = client.get("/api/v1/households").json()["data"]["items"][0]
        assert item["member_count"] == 0

    def test_member_not_found_and_wrong_household(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = self._create_household(client, village.id)
        other_hid = self._create_household(client, village.id)
        member = _mk_member(db, other_hid, name="别人")
        assert client.put(
            f"/api/v1/households/{hid}/members/9999", json={"name": "x"}
        ).status_code == 404
        assert client.put(
            f"/api/v1/households/{hid}/members/{member.id}", json={"name": "x"}
        ).status_code == 404
        assert client.delete(f"/api/v1/households/{hid}/members/9999").status_code == 404

    def test_member_cross_scope_not_found(self, api_client):
        client, db, state = api_client
        village = _mk_village(db, org_id=1)
        h = _mk_household(db, village.id)
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert client.get(f"/api/v1/households/{h.id}/members").status_code == 404
        assert (
            client.post(
                f"/api/v1/households/{h.id}/members", json={"name": "甲"}
            ).status_code
            == 404
        )


# ════════════════════════════════════════════════════════════════════
#  预警规则
# ════════════════════════════════════════════════════════════════════

def _household_stub(**kw):
    base = dict(
        household_type="monitor",
        last_visit_date=None,
        has_serious_illness=False,
        housing_safe=True,
        risk_type="none",
    )
    base.update(kw)
    return SimpleNamespace(**base)


class TestAlertRules:
    REFERENCE = date(2026, 6, 1)

    def test_never_visited(self):
        from app.services.household_service import evaluate_household_alerts

        alerts = evaluate_household_alerts(_household_stub(), today=self.REFERENCE)
        assert [alert["code"] for alert in alerts] == ["no_visit"]
        assert alerts[0]["level"] == "medium"

    def test_general_household_never_visited_no_alert(self):
        from app.services.household_service import evaluate_household_alerts

        alerts = evaluate_household_alerts(
            _household_stub(household_type="general"), today=self.REFERENCE
        )
        assert alerts == []

    def test_boundary_90_days_no_alert(self):
        from app.services.household_service import evaluate_household_alerts

        visited = self.REFERENCE - timedelta(days=90)
        alerts = evaluate_household_alerts(
            _household_stub(last_visit_date=visited), today=self.REFERENCE
        )
        assert alerts == []

    def test_boundary_91_days_alert(self):
        from app.services.household_service import evaluate_household_alerts

        visited = self.REFERENCE - timedelta(days=91)
        alerts = evaluate_household_alerts(
            _household_stub(last_visit_date=visited), today=self.REFERENCE
        )
        assert [alert["code"] for alert in alerts] == ["no_visit"]

    def test_high_rules_combined(self):
        from app.services.household_service import evaluate_household_alerts

        alerts = evaluate_household_alerts(
            _household_stub(
                household_type="general",
                has_serious_illness=True,
                housing_safe=False,
                risk_type="multiple",
            ),
            today=self.REFERENCE,
        )
        assert [alert["code"] for alert in alerts] == [
            "serious_illness",
            "housing_unsafe",
            "risk_type",
        ]
        assert all(alert["level"] == "high" for alert in alerts)

    @pytest.mark.parametrize("risk_type", ["sudden", "multiple"])
    def test_high_risk_types(self, risk_type):
        from app.services.household_service import evaluate_household_alerts

        alerts = evaluate_household_alerts(
            _household_stub(household_type="general", risk_type=risk_type),
            today=self.REFERENCE,
        )
        assert [alert["code"] for alert in alerts] == ["risk_type"]

    def test_household_type_poverty_uses_same_no_visit_rule(self):
        from app.services.household_service import evaluate_household_alerts

        alerts = evaluate_household_alerts(
            _household_stub(household_type="poverty"), today=self.REFERENCE
        )
        assert [alert["code"] for alert in alerts] == ["no_visit"]

    def test_alerts_api_rule_and_open_records_and_level_filter(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        risky = _mk_household(
            db,
            village.id,
            household_name="风险户",
            household_type="monitor",
            has_serious_illness=True,
        )
        _mk_risk_record(
            db, risky.id, risk_type="income_drop", risk_level="medium", reason="收入骤减"
        )
        _mk_risk_record(db, risky.id, risk_type="illness", risk_level="high", status="closed")
        _mk_household(db, village.id, household_name="平稳户", household_type="general")

        data = client.get("/api/v1/households/alerts").json()["data"]
        assert data["total"] == len(data["items"])
        codes = {item["code"] for item in data["items"]}
        assert "serious_illness" in codes
        assert "no_visit" in codes
        assert "income_drop" in codes
        assert "illness" not in codes  # 已关闭的风险记录不进入预警清单
        # 高等级优先排序
        levels = [item["level"] for item in data["items"]]
        assert levels == sorted(levels, key=lambda value: {"high": 0, "medium": 1, "low": 2}[value])

        high_only = client.get("/api/v1/households/alerts", params={"level": "high"}).json()[
            "data"
        ]
        assert high_only["total"] >= 1
        assert all(item["level"] == "high" for item in high_only["items"])

    def test_alerts_skip_orphan_records_of_soft_deleted_household(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        h = _mk_household(db, village.id, household_type="general")
        _mk_risk_record(db, h.id, risk_type="income_drop", reason="待关闭")
        client.delete(f"/api/v1/households/{h.id}")
        data = client.get("/api/v1/households/alerts").json()["data"]
        assert data["total"] == 0

    def test_list_alerts_empty(self, api_client):
        client, _db, _state = api_client
        data = client.get("/api/v1/households/alerts").json()["data"]
        assert data == {"items": [], "total": 0, "page": 1, "page_size": 1}


# ════════════════════════════════════════════════════════════════════
#  风险记录
# ════════════════════════════════════════════════════════════════════

class TestRiskRecords:
    def test_create_list_and_close(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        hid = _post(client, _base_payload(village.id, risk_level="high")).json()["data"]["id"]

        created = client.post(
            f"/api/v1/households/{hid}/risk-records",
            json={"risk_type": "income_drop", "reason": "收入下降 30%"},
        )
        assert created.status_code == 200
        record = created.json()["data"]
        assert record["status"] == "open"
        assert record["risk_level"] == "high"  # 缺省继承户风险等级
        assert record["detected_at"]

        explicit = client.post(
            f"/api/v1/households/{hid}/risk-records",
            json={
                "risk_type": "housing",
                "risk_level": "low",
                "reason": "住房隐患",
                "detected_at": "2026-01-02T03:04:05+00:00",
            },
        ).json()["data"]
        assert explicit["risk_level"] == "low"
        assert explicit["detected_at"].startswith("2026-01-02")

        listed = client.get(f"/api/v1/households/{hid}/risk-records").json()["data"]
        assert listed["total"] == 2
        # 按发现时间倒序：显式 2026-01-02 与创建时 now（now > 2026-01？取决于运行时间，改为按集合断言）
        codes = {item["risk_type"] for item in listed["items"]}
        assert codes == {"income_drop", "housing"}

        closed = client.put(
            f"/api/v1/households/risk-records/{record['id']}/close",
            json={"closed_note": "已落实帮扶措施"},
        )
        assert closed.status_code == 200
        assert closed.json()["data"]["status"] == "closed"
        assert closed.json()["data"]["closed_at"]
        assert closed.json()["data"]["closed_note"] == "已落实帮扶措施"

        # 重复关闭 → 400
        again = client.put(
            f"/api/v1/households/risk-records/{record['id']}/close", json={}
        )
        assert again.status_code == 400
        assert "已关闭" in again.json()["message"]

    def test_close_not_found_and_cross_scope(self, api_client):
        client, db, state = api_client
        village = _mk_village(db, org_id=1)
        h = _mk_household(db, village.id)
        record = _mk_risk_record(db, h.id)
        assert (
            client.put("/api/v1/households/risk-records/9999/close", json={}).status_code
            == 404
        )
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert (
            client.put(
                f"/api/v1/households/risk-records/{record.id}/close", json={}
            ).status_code
            == 404
        )

    def test_create_risk_record_for_cross_scope_household(self, api_client):
        client, db, state = api_client
        village = _mk_village(db, org_id=1)
        h = _mk_household(db, village.id)
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        resp = client.post(
            f"/api/v1/households/{h.id}/risk-records", json={"risk_type": "income_drop"}
        )
        assert resp.status_code == 404


# ════════════════════════════════════════════════════════════════════
#  汇总回填
# ════════════════════════════════════════════════════════════════════

class TestRollup:
    def _seed_village(self, db):
        village = _mk_village(db, name="幸福村")
        poverty = _mk_household(
            db, village.id, household_name="脱贫户", household_type="poverty"
        )
        monitor = _mk_household(
            db, village.id, household_name="监测户", household_type="monitor"
        )
        _mk_household(db, village.id, household_name="一般户")
        _mk_member(db, poverty.id)
        _mk_member(db, poverty.id)
        _mk_member(db, monitor.id)
        deleted_member = _mk_member(db, monitor.id, name="已删成员")
        deleted_member.is_active = False
        db.commit()
        deleted_household = _mk_household(db, village.id, household_name="已删户")
        deleted_household.is_active = False
        db.commit()
        _mk_member(db, deleted_household.id)
        return village

    def test_rollup_creates_and_writes_back(self, api_client):
        client, db, _state = api_client
        village = self._seed_village(db)
        resp = client.post(
            "/api/v1/households/rollup", params={"village_id": village.id, "year": 2025}
        )
        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["village_name"] == "幸福村"
        assert data["year"] == 2025
        assert data["total_households"] == 3
        assert data["total_population"] == 3  # 已删成员 / 已删户不计入
        assert data["poverty_households"] == 1
        assert data["monitoring_households"] == 1
        assert data["created"] is True
        assert data["written_fields"] == [
            "total_households",
            "total_population",
            "poverty_households",
            "unstable_poverty_households",
        ]

        from app.models.supported_village import VillagePopulation

        row = (
            db.query(VillagePopulation)
            .filter(
                VillagePopulation.supported_village_id == village.id,
                VillagePopulation.year == 2025,
            )
            .first()
        )
        assert row is not None
        assert row.total_households == 3
        assert row.total_population == 3
        assert row.poverty_households == 1
        assert row.unstable_poverty_households == 1

    def test_rollup_idempotent_and_updates_existing_row(self, api_client):
        client, db, _state = api_client
        village = self._seed_village(db)
        _mk_population(db, village.id, 2025, total_households=99, resident_population=88)
        first = client.post(
            "/api/v1/households/rollup", params={"village_id": village.id, "year": 2025}
        ).json()["data"]
        assert first["created"] is False
        second = client.post(
            "/api/v1/households/rollup", params={"village_id": village.id, "year": 2025}
        ).json()["data"]
        assert second["created"] is False
        assert second["total_households"] == first["total_households"] == 3

        from app.models.supported_village import VillagePopulation

        rows = (
            db.query(VillagePopulation)
            .filter(VillagePopulation.supported_village_id == village.id)
            .all()
        )
        assert len(rows) == 1  # 幂等 upsert，不重复插行
        assert rows[0].resident_population == 88  # 不覆盖未参与汇总的既有字段

    def test_rollup_default_year_is_current(self, api_client):
        client, db, _state = api_client
        village = self._seed_village(db)
        current_year = datetime.now(timezone.utc).year
        data = client.post(
            "/api/v1/households/rollup", params={"village_id": village.id}
        ).json()["data"]
        assert data["year"] == current_year

    def test_rollup_empty_village_writes_zeros(self, api_client):
        client, db, _state = api_client
        village = _mk_village(db)
        data = client.post(
            "/api/v1/households/rollup", params={"village_id": village.id, "year": 2024}
        ).json()["data"]
        assert data["total_households"] == 0
        assert data["total_population"] == 0
        assert data["poverty_households"] == 0
        assert data["monitoring_households"] == 0

    def test_rollup_village_not_found_and_cross_scope(self, api_client):
        client, db, state = api_client
        village = _mk_village(db, org_id=1)
        assert (
            client.post(
                "/api/v1/households/rollup", params={"village_id": 9999, "year": 2025}
            ).status_code
            == 404
        )
        _as(state, id=9, role="user", organization_id=2, data_scope="org")
        assert (
            client.post(
                "/api/v1/households/rollup", params={"village_id": village.id, "year": 2025}
            ).status_code
            == 404
        )


# ════════════════════════════════════════════════════════════════════
#  导出
# ════════════════════════════════════════════════════════════════════

class TestExport:
    def _seed(self, db):
        village = _mk_village(db, name="幸福村")
        household = _mk_household(
            db,
            village.id,
            household_name="张三",
            household_code="H-001",
            household_type="monitor",
            risk_level="high",
            helper_name="王干事",
            helper_phone="13800138000",
            annual_income=12000,
            migrant_workers=2,
            has_serious_illness=True,
            housing_safe=False,
            last_visit_date=date(2026, 1, 1),
        )
        _mk_member(db, household.id)
        _mk_household(db, village.id, household_name="李四", risk_level="low")
        return village

    def test_export_xlsx(self, api_client):
        import io

        from openpyxl import load_workbook

        client, db, _state = api_client
        self._seed(db)
        resp = client.get("/api/v1/households/export", params={"format": "xlsx"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "households_" in resp.headers["content-disposition"]
        assert ".xlsx" in resp.headers["content-disposition"]
        workbook = load_workbook(io.BytesIO(resp.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0][0] == "户编号"
        assert rows[0][2] == "帮扶村"
        assert len(rows) == 3  # 表头 + 2 户
        by_name = {row[1]: row for row in rows[1:]}
        assert by_name["张三"][2] == "幸福村"
        assert by_name["张三"][10] == "是"  # 大病 → 是
        assert by_name["张三"][11] == "否"  # 住房安全 → 否
        assert "未走访" in by_name["张三"][14]  # 预警文案

    def test_export_csv_with_filters(self, api_client):
        client, db, _state = api_client
        self._seed(db)
        resp = client.get(
            "/api/v1/households/export",
            params={"format": "csv", "risk_level": "high"},
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/csv")
        assert ".csv" in resp.headers["content-disposition"]
        text = resp.content.decode("utf-8-sig")
        lines = text.splitlines()
        assert lines[0].startswith("户编号,户主姓名,帮扶村")
        assert len(lines) == 2  # 表头 + 仅高风险户
        assert "张三" in lines[1]

    def test_export_village_filter_and_include_deleted(self, api_client):
        client, db, state = api_client
        village = self._seed(db)
        deleted = _mk_household(db, village.id, household_name="已删户")
        deleted.is_active = False
        db.commit()
        resp = client.get(
            "/api/v1/households/export",
            params={"format": "csv", "village_id": village.id, "include_deleted": "true"},
        )
        assert resp.status_code == 200
        assert "已删户" in resp.content.decode("utf-8-sig")
        _as(state, id=9, role="user", organization_id=1, data_scope="org")
        resp = client.get(
            "/api/v1/households/export",
            params={"format": "csv", "include_deleted": "true"},
        )
        assert "已删户" not in resp.content.decode("utf-8-sig")

    def test_export_invalid_format(self, api_client):
        client, db, _state = api_client
        _mk_village(db)
        resp = client.get("/api/v1/households/export", params={"format": "pdf"})
        assert resp.status_code == 400
        assert "不支持的导出格式" in resp.json()["message"]


# ════════════════════════════════════════════════════════════════════
#  服务层直接调用（分支补盲：include_deleted / 空集合 / 归一化）
# ════════════════════════════════════════════════════════════════════

class TestServiceInternals:
    def test_base_queries_include_deleted_branch(self, api_client):
        from app.core.data_permission import OrgScopeFilter
        from app.services import household_service

        _client, db, _state = api_client
        village = _mk_village(db)
        household = _mk_household(db, village.id)
        member = _mk_member(db, household.id)
        record = _mk_risk_record(db, household.id)
        household.is_active = False
        member.is_active = False
        record.is_active = False
        db.commit()
        scope = OrgScopeFilter(is_admin=True)

        assert household_service.household_base_query(db, scope).count() == 0
        assert household_service.household_base_query(db, scope, include_deleted=True).count() == 1
        assert (
            household_service.member_base_query(db, scope, household.id).count() == 0
        )
        assert (
            household_service.member_base_query(
                db, scope, household.id, include_deleted=True
            ).count()
            == 1
        )
        assert (
            household_service.risk_record_base_query(db, scope).count() == 0
        )
        assert (
            household_service.risk_record_base_query(
                db, scope, include_deleted=True
            ).count()
            == 1
        )

    def test_empty_batch_helpers_and_enrich(self, api_client):
        from app.services import household_service

        _client, db, _state = api_client
        assert household_service.enrich_households(db, []) == []
        assert household_service._village_name_map(db, set()) == {}
        assert household_service._member_count_map(db, []) == {}
        assert household_service._today() == datetime.now(timezone.utc).date()

    def test_export_value_normalization(self):
        from app.services.household_service import _export_value

        assert _export_value(None) == ""
        assert _export_value(True) == "是"
        assert _export_value(False) == "否"
        assert _export_value([{"message": "预警一"}, {"message": "预警二"}]) == "预警一；预警二"
        assert _export_value(12.5) == 12.5

    def test_build_bytes_direct(self):
        from openpyxl import load_workbook

        from app.services.household_service import build_csv_bytes, build_xlsx_bytes

        headers = ["列一", "列二"]
        rows = [["甲", 1], ["乙", 2]]
        csv_bytes = build_csv_bytes(headers, rows)
        assert csv_bytes.startswith(b"\xef\xbb\xbf")
        assert "甲" in csv_bytes.decode("utf-8-sig")

        xlsx_bytes = build_xlsx_bytes(headers, rows)
        workbook = load_workbook(__import__("io").BytesIO(xlsx_bytes))
        assert list(workbook.active.iter_rows(values_only=True))[0] == ("列一", "列二")

    def test_serializers(self, api_client):
        from app.services import household_service

        _client, db, _state = api_client
        village = _mk_village(db)
        household = _mk_household(db, village.id)
        member = _mk_member(db, household.id)
        record = _mk_risk_record(db, household.id)
        assert household_service.serialize_household(household)["household_name"] == "张三"
        assert household_service.serialize_member(member)["name"] == "李四"
        assert household_service.serialize_risk_record(record)["risk_type"] == "income_drop"

    def test_household_export_data_direct_with_include_deleted(self, api_client):
        from app.core.data_permission import OrgScopeFilter
        from app.services import household_service

        _client, db, _state = api_client
        village = _mk_village(db)
        household = _mk_household(db, village.id)
        household.is_active = False
        db.commit()
        scope = OrgScopeFilter(is_admin=True)
        headers, rows = household_service.household_export_data(db, scope, include_deleted=True)
        assert headers[0] == "户编号"
        assert len(rows) == 1
