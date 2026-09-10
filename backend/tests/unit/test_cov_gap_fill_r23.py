"""覆盖率补充测试：`report_export_service` 三处未覆盖分支 + `security` 用户不存在分支。

全仓可覆盖集门禁（100%）此前差 4 行，均在他人本轮改动的新增/改动代码里，且没有
测试触达；本文件只补测试，不改动被测源码：

- `report_export_service.py:390`：Word 导出时表格 payload 的 `headers` 为空 → 跳过建表；
- `report_export_service.py:447`：`_style_paragraph_runs(size=...)` 的显式字号分支；
- `report_export_service.py:545`：PDF 导出时 `headers` 为空 → 跳过建表；
- `core/security.py:338`：`get_current_user` 查不到用户 → 401「用户不存在」。
"""

import inspect
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from fastapi import HTTPException

import app.core.security as sec
from app.services.report_export_service import ReportExportService


class TestWordEmptyHeadersTable:
    def test_empty_headers_table_is_skipped(self):
        svc = ReportExportService()
        data = {
            "title": "空表头报表",
            "sections": [{"title": "段", "paragraphs": ["说明"], "table": {"headers": [], "rows": [["a"]]}}],
        }
        out = svc.export_word("summary", data)
        assert out[:2] == b"PK"  # docx 是 zip 容器

    def test_table_with_headers_still_rendered(self):
        svc = ReportExportService()
        data = {
            "title": "正常报表",
            "sections": [{"title": "段", "paragraphs": [], "table": {"headers": ["A"], "rows": [["1"]]}}],
        }
        assert svc.export_word("summary", data)[:2] == b"PK"


class TestStyleParagraphRunsSize:
    def test_explicit_size_branch(self):
        from docx import Document

        doc = Document()
        p = doc.add_paragraph("文本")
        ReportExportService._style_paragraph_runs(p, bold=True, color="FFFFFF", center=True, size=9)
        assert p.runs and p.runs[0].font.size is not None

    def test_default_size_branch(self):
        from docx import Document

        doc = Document()
        p = doc.add_paragraph("文本")
        ReportExportService._style_paragraph_runs(p)
        assert p.runs[0].font.bold is False


class TestPdfEmptyHeadersTable:
    def test_empty_headers_table_is_skipped(self):
        svc = ReportExportService()
        data = {
            "title": "空表头 PDF",
            "year": 2026,
            "sections": [{"title": "段", "paragraphs": ["说明"], "table": {"headers": [], "rows": [["a"]]}}],
        }
        assert svc.export_pdf("summary", data).startswith(b"%PDF")


class TestGetCurrentUserUserMissing:
    async def test_missing_user_returns_401(self, monkeypatch):
        db = MagicMock()
        db.query.return_value.filter.return_value.first.return_value = None
        monkeypatch.setattr(
            sec, "decode_token", MagicMock(return_value={"sub": "ghost", "type": "access"})
        )

        credentials = SimpleNamespace(credentials="tok")
        kwargs = {"credentials": credentials}
        if "db" in inspect.signature(sec.get_current_user).parameters:
            kwargs["db"] = db
        else:  # pragma: no cover - 旧签名兼容分支（当前源码已是注入式）
            monkeypatch.setattr("app.core.database.SessionLocal", MagicMock(return_value=db))

        with pytest.raises(HTTPException) as exc:
            await sec.get_current_user(**kwargs)

        assert exc.value.status_code == 401
        assert exc.value.detail == "用户不存在"
