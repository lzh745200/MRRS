# -*- coding: utf-8 -*-
# 固化自 2026-09 深度探测循环（r4 轮）。
# 运行：python backend/tests/probe/probe_r4_policy_school_worklog_2fa.py
#      或一次跑全部：python backend/tests/probe/run_all.py
# 隔离契约：BUMOFU_BACKEND_DIR_OVERRIDE + DATABASE_URL 均指向本次运行临时目录，
#           不触碰真实 backend/data；退出码 0=全过，1=有失败。
"""第 4 轮模块深探：政策/学校/项目/工作日志/通知/2FA/资料/报表订阅/模板上传确认（真实 HTTP）"""
import io
import os
import sys
import tempfile

TMP = tempfile.mkdtemp(prefix="probe_r4_")
os.chdir(TMP)
os.environ["BUMOFU_BACKEND_DIR_OVERRIDE"] = TMP
os.environ["DATABASE_URL"] = f"sqlite:///{TMP}/probe.db"
os.environ["ENVIRONMENT"] = "test"
os.environ["SECRET_KEY"] = "x" * 40
os.environ["UPLOAD_DIR"] = f"{TMP}/uploads"
os.environ["BACKUP_ENABLED"] = "false"

BACKEND_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BACKEND_ROOT)

from openpyxl import Workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("PASS" if cond else "FAIL"), name, "" if cond else str(detail)[:240])


def unwrap(j):
    return j.get("data") if isinstance(j, dict) and j.get("code") == 200 and isinstance(j.get("data"), (dict, list)) else j


def safe_json(r):
    try:
        return r.json()
    except Exception:
        return {"_non_json": True}


from app.main import app  # noqa: E402

with TestClient(app, raise_server_exceptions=False) as client:
    r = client.get("/api/v1/auth/csrf-token")
    token = (r.json().get("data") or {}).get("csrf_token")
    r = client.post("/api/v1/auth/login", json={"username": "admin", "password": "Admin@2026"},
                    headers={"X-CSRF-Token": token})
    at = (r.json().get("data") or {}).get("access_token")
    check("管理员登录", bool(at))
    AUTH = {"Authorization": f"Bearer {at}", "X-CSRF-Token": token}

    # ═══ 链路 1：政策法规 ═══
    r = client.post("/api/v1/policies",
                    json={"title": "探针政策-产业帮扶办法", "content": "第一条 总则", "category": "产业",
                          "status": "published"}, headers=AUTH)
    pj = r.json()
    pol_id = (unwrap(pj) or {}).get("id") or pj.get("id")
    check("创建政策", bool(pol_id), pj)
    r = client.get("/api/v1/policies", headers=AUTH)
    check("政策列表", r.status_code == 200, r.status_code)
    r = client.put(f"/api/v1/policies/{pol_id}",
                   json={"title": "探针政策-改", "content": "第二条 修订"}, headers=AUTH)
    check("更新政策", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/policies/categories", headers=AUTH)
    check("政策分类", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/policies/export/excel", headers=AUTH)
    check("政策导出Excel", r.status_code == 200 and len(r.content) > 300, (r.status_code, len(r.content)))
    r = client.delete(f"/api/v1/policies/{pol_id}", headers=AUTH)
    check("删除政策", r.status_code == 200, r.status_code)

    # ═══ 链路 2：学校 ═══
    r = client.post("/api/v1/schools",
                    json={"name": "探针希望小学", "code": "SCH-PROBE-01", "school_type": "primary",
                          "school_level": "county", "province": "贵州省"},
                    headers=AUTH)
    sj = r.json()
    sch_id = (unwrap(sj) or {}).get("id")
    check("创建学校", bool(sch_id), sj)
    r = client.get("/api/v1/schools", headers=AUTH)
    check("学校列表", r.status_code == 200, r.status_code)
    r = client.put(f"/api/v1/schools/{sch_id}", json={"name": "探针希望小学-改"}, headers=AUTH)
    check("更新学校", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/schools/statistics", headers=AUTH)
    check("学校统计", r.status_code == 200, r.status_code)
    r = client.delete(f"/api/v1/schools/{sch_id}", headers=AUTH)
    check("删除学校", r.status_code == 200, r.status_code)

    # ═══ 链路 3：项目 ═══
    r = client.post("/api/v1/projects",
                    json={"name": "探针项目-饮水工程", "type": "基础设施", "budget": 200000},
                    headers=AUTH)
    prj = r.json()
    proj_id = (unwrap(prj) or {}).get("id") or prj.get("id")
    check("创建项目", bool(proj_id), prj)
    r = client.get(f"/api/v1/projects/{proj_id}", headers=AUTH)
    check("项目详情", r.status_code == 200, r.status_code)
    r = client.put(f"/api/v1/projects/{proj_id}", json={"name": "探针项目-改"}, headers=AUTH)
    check("更新项目", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/projects/stats", headers=AUTH)
    check("项目统计", r.status_code == 200, r.status_code)
    r = client.delete(f"/api/v1/projects/{proj_id}", headers=AUTH)
    check("删除项目", r.status_code == 200, r.status_code)

    # ═══ 链路 4：工作日志 ═══
    r = client.post("/api/v1/work-logs",
                    json={"title": "探针日志-入户走访", "content": "走访 5 户，收集诉求 3 条",
                          "work_date": "2026-09-05", "log_type": "走访"},
                    headers=AUTH)
    wj = r.json()
    wl_id = (unwrap(wj) or {}).get("id") or wj.get("id")
    check("创建工作日志", bool(wl_id), wj)
    r = client.get("/api/v1/work-logs", headers=AUTH)
    check("工作日志列表", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/work-logs/calendar?year=2026&month=9", headers=AUTH)
    check("工作日历", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/work-logs/monthly-summary?year=2026&month=9", headers=AUTH)
    check("月度总结", r.status_code == 200, r.status_code)
    r = client.put(f"/api/v1/work-logs/{wl_id}", json={"content": "走访 6 户"}, headers=AUTH)
    check("更新日志", r.status_code == 200, r.status_code)
    r = client.delete(f"/api/v1/work-logs/{wl_id}", headers=AUTH)
    check("删除日志", r.status_code == 200, r.status_code)

    # ═══ 链路 5：通知设置 ═══
    r = client.get("/api/v1/notifications/preferences", headers=AUTH)
    check("通知偏好读取", r.status_code == 200, r.status_code)
    r = client.get("/api/v1/notifications", headers=AUTH)
    check("通知列表", r.status_code in (200, 404), r.status_code)

    # ═══ 链路 6：2FA 设置 ═══
    r = client.post("/api/v1/two-factor/enable", headers=AUTH)
    ej = r.json()
    ed = unwrap(ej) or ej
    secret = (ed or {}).get("secret")
    check("启用2FA(返回密钥)", r.status_code == 200 and bool(secret), ej)
    r = client.get("/api/v1/two-factor/status", headers=AUTH)
    st = (unwrap(r.json()) or r.json())
    enabled_now = st.get("enabled") if isinstance(st, dict) else None
    check("2FA状态查询", r.status_code == 200, r.status_code)
    # 错误验证码 verify → 400
    r = client.post("/api/v1/two-factor/verify", json={"token": "000000"}, headers=AUTH)
    check("错误验证码被拒(400)", r.status_code == 400, (r.status_code, str(r.json())[:120]))
    # 关闭 2FA（恢复环境；带验证码或密码按契约）
    r = client.post("/api/v1/two-factor/disable", json={"code": "000000", "password": "Admin@2026"},
                    headers=AUTH)
    check("关闭2FA", r.status_code in (200, 400), (r.status_code, str(r.json())[:120]))

    # ═══ 链路 7：个人资料 ═══
    r = client.get("/api/v1/users/me", headers=AUTH)
    me = unwrap(r.json()) or {}
    my_id = me.get("id")
    check("当前用户信息", r.status_code == 200 and bool(my_id), r.json())
    r = client.put("/api/v1/users/me/profile",
                   json={"full_name": "系统管理员-探针", "email": "admin@example.com"}, headers=AUTH)
    check("更新个人资料", r.status_code == 200, str(r.json())[:140])
    r = client.get("/api/v1/users/me", headers=AUTH)
    check("资料回读", (unwrap(r.json()) or {}).get("name") == "系统管理员-探针",
          (unwrap(r.json()) or {}).get("full_name"))

    # ═══ 链路 8：报表订阅 ═══
    r = client.post("/api/v1/reports/subscriptions",
                    json={"name": "探针周报订阅", "report_type": "village_yearly", "format": "xlsx",
                          "frequency": "weekly"}, headers=AUTH)
    bj = r.json()
    sub_id = (unwrap(bj) or {}).get("id")
    check("创建报表订阅", bool(sub_id), bj)
    r = client.get("/api/v1/reports/subscriptions", headers=AUTH)
    check("订阅列表", r.status_code == 200, r.status_code)
    r = client.get(f"/api/v1/reports/subscriptions/{sub_id}", headers=AUTH)
    check("订阅详情", r.status_code == 200, r.status_code)

    # ═══ 链路 9：模板上传确认（preview→confirm 两段式）═══
    r = client.post("/api/v1/report-templates",
                    json={"name": "探针上传模板", "type": "import", "module": "village",
                          "fields": "village_name,province"}, headers=AUTH)
    tj = r.json()
    tpl_id = (unwrap(tj) or {}).get("id") or tj.get("id") or tj.get("template_id")
    check("建导入模板", bool(tpl_id), tj)

    # 构造已填写的模板 Excel：表头行2 + 数据行3（下载端点约定：1 标题 / 2 表头 / 3+ 数据）
    r = client.get(f"/api/v1/report-templates/{tpl_id}/download", headers=AUTH)
    check("下载空白模板", r.status_code == 200, r.status_code)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "探针上传模板"
    ws["A2"], ws["B2"] = "village_name", "province"
    ws["A3"], ws["B3"] = "探针模板村A", "贵州省"
    ws["A4"], ws["B4"] = "探针模板村B", "贵州省"
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post(f"/api/v1/report-templates/{tpl_id}/upload?mode=preview",
                    files={"file": ("filled.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    headers=AUTH)
    pvj = r.json()
    pvd = unwrap(pvj) or pvj
    check("模板预览解析", r.status_code == 200 and str(pvd).find("探针模板村A") != -1, str(pvj)[:220])

    r = client.post(f"/api/v1/report-templates/{tpl_id}/upload?mode=confirm",
                    files={"file": ("filled.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    headers=AUTH)
    cfj = r.json()
    cfd = unwrap(cfj) or cfj
    check("模板确认导入", r.status_code == 200, str(cfj)[:220])

    # 确认后数据真实落库
    r = client.get("/api/v1/supported-villages", params={"keyword": "探针模板村A"}, headers=AUTH)
    items = unwrap(r.json()) or []
    if isinstance(items, dict):
        items = items.get("items") or []
    landed = any((it.get("village_name") or it.get("name")) == "探针模板村A" for it in items)
    check("确认导入真实落库", landed, str(items)[:200])

print(f"\n===== PASS {len(PASS)} / FAIL {len(FAIL)} =====")
for n, dd in FAIL:
    print("  FAIL:", n, "|", str(dd)[:220])
sys.exit(1 if FAIL else 0)
